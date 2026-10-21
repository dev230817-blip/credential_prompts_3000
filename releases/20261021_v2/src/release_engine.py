"""Offline option-A release: supplied values fixed, model chooses layout and fixed prose."""
from __future__ import annotations
import collections,csv,datetime,hashlib,io,json,pathlib,zipfile,re

VERSION='credential-knowledge-guided-20261021-v2'
RELEASE='20261021_v2'
SEMANTICS='same_supplied_core_values; layout_and_fixed_prose_may_vary'
COLUMNS=['sample_id','direction','category_l1','category_l2','variant','aspect_ratio','prompt','group_id','target_texts','tamper_types','tamper_details','extra_json']
PHYSICAL=('scan','front_photo','oblique_photo')
ELECTRONIC=('digital_flat','screen_front_photo','screen_oblique_photo')

def sha(data):return hashlib.sha256(data).hexdigest()
def canonical(v):return json.dumps(v,ensure_ascii=False,sort_keys=True,separators=(',',':'))
def read_json(p):return json.loads(p.read_text(encoding='utf-8'))
def read_jl(p):return [json.loads(s) for s in p.read_text(encoding='utf-8').splitlines() if s.strip()]
def text_file(p,s):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(s,encoding='utf-8',newline='\n')
def dump_json(p,v):text_file(p,json.dumps(v,ensure_ascii=False,indent=2)+'\n')
def dump_jl(p,rows):text_file(p,''.join(canonical(r)+'\n' for r in rows))
def dump_csv(p,rows,fields=None):
    p.parent.mkdir(parents=True,exist_ok=True)
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fields or list(rows[0]));w.writeheader();w.writerows(rows)
def read_csv(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))

def core_payload(b):
    return {k:b[k] for k in ['category_entry_id','core_fields','material_kind','portrait_instruction']}

def base_prompt(b,policy):
    context='；'.join(f'{k}：{v}' for k,v in [('国家或地区',b['country']),('地区',b['region'])] if v)
    vals='\n'.join(f'{f["label"]}：{f["value"]}' for f in b['core_fields'])
    portrait=('\n'+b['portrait_instruction']) if b['portrait_instruction'] else ''
    return ('生成“'+b['category_name']+'”的一张合成文书图像，使用下方完整给定的虚构样例取值。\n\n'
            +'【对象与样式】\n'+(context+'。' if context else '')+policy['knowledge_instruction']+'\n'+policy['context_policy']
            +'\n\n【给定核心取值】\n'+vals+'\n\n【内容边界】\n'+policy['values_policy']+'\n'+policy['unprovided_policy']+portrait
            +'\n\n【生成要求】\n'+policy['output_policy'])

def render(b,mode,config):
    cap=config['modes'][mode]['text']
    canvas=f'输出画布宽:高={b["canvas_aspect_ratio"]}。画布比例只规定输出画幅，文档本体的横竖方向与比例由其样式决定，以留白容纳完整文档，不拉伸文档填满画布。'
    return b['base_prompt']+'\n\n【本条成像要求】\n'+cap+'\n'+canvas+'\n\n直接生成图像，不输出设计说明。'

def generate(repo):
    bases=read_jl(repo/'data/base_values_1430.jsonl');config=read_json(repo/'config/imaging_rules.json');rows=[]
    for b in bases:
        for mode in ELECTRONIC if b['material_kind']=='electronic' else PHYSICAL:
            prompt=render(b,mode,config)
            r={k:b[k] for k in ['category_entry_id','category_l1_id','category_l1','category_l2_id','category_l2','category_name','country','region','document_family','family_label','material_kind','core_fields','target_texts','expected_core_values','core_values_sha256','core_value_group_id','source_base_sample_id','source_prompt_sha256','source_base_version','source_release_version','source_parent_samples','source_ids','evidence_status','layout_status','canvas_aspect_ratio','portrait_instruction']}
            r.update(sample_id=f'CRD-20261021-v2-{len(rows)+1:05d}',version=VERSION,release_version=RELEASE,direction='credential',variant='clean',group_id='',tamper_types='',tamper_details='',capture_mode=mode,capture_label=config['modes'][mode]['label'],aspect_ratio=b['canvas_aspect_ratio'],document_aspect_ratio=None,document_aspect_ratio_policy='model_knowledge; not_locked',content_group_id=b['core_value_group_id'],content_sha256=b['core_values_sha256'],content_hash_scope='supplied_core_values_and_subject_only; not_full_page',group_semantics=SEMANTICS,target_texts_scope='supplied_core_values_only',layout_locked=False,full_visible_text_locked=False,additional_fixed_prose_allowed=True,prompt=prompt,prompt_sha256=sha(prompt.encode()),new_release_image_generated=False,new_release_image_accepted=False,new_prompt_individually_approved=False)
            rows.append(r)
    return bases,rows

def team_rows(rows):
    result=[]
    for r in rows:
        x={k:r[k] for k in COLUMNS if k!='extra_json'}
        x['extra_json']=canonical({k:v for k,v in r.items() if k not in set(COLUMNS)})
        result.append(x)
    return result

def xlsx(p,rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    w=Workbook();s=w.active;s.title='4290条A口径'
    keys=['sample_id','category_entry_id','category_l1','category_l2','category_name','country','region','material_kind','capture_mode','aspect_ratio','core_value_group_id','group_semantics','prompt']
    s.append(keys)
    for r in rows:s.append([r[k] for k in keys])
    s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
    for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='274D72')
    for col,width in zip('ABCDEFGHIJKLM',[32,19,20,32,55,16,18,15,24,14,38,55,110]):s.column_dimensions[col].width=width
    for row in s.iter_rows(min_row=2):row[-1].alignment=Alignment(wrap_text=True,vertical='top')
    fixed=datetime.datetime(2026,10,21);w.properties.creator='credential-prompt-project';w.properties.lastModifiedBy='credential-prompt-project';w.properties.created=fixed;w.properties.modified=fixed
    memory=io.BytesIO();w.save(memory)
    with zipfile.ZipFile(io.BytesIO(memory.getvalue())) as source,zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED) as out:
        for name in sorted(source.namelist()):
            data=source.read(name)
            if name=='docProps/core.xml':data=re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)',rb'\g<1>2026-10-21T00:00:00Z\2',data)
            info=zipfile.ZipInfo(name,(2026,10,21,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;info.external_attr=0o600<<16;out.writestr(info,data)

def export(repo,rows):
    dump_jl(repo/'data/prompts.jsonl',rows)
    dump_csv(repo/'data/credential_prompts_4290_unified.csv',team_rows(rows),COLUMNS)
    text_file(repo/'data/prompts.txt','\n\n'.join(f'===== {r["sample_id"]} | {r["category_name"]} | {r["capture_label"]} =====\n{r["prompt"]}' for r in rows)+'\n')
    xlsx(repo/'data/prompts_filter.xlsx',rows)

def stats(bases,rows):
    return dict(version=VERSION,policy='A',group_semantics=SEMANTICS,categories=len(bases),prompts=len(rows),category_l1=dict(collections.Counter(b['category_l1'] for b in bases)),capture_modes=dict(collections.Counter(r['capture_mode'] for r in rows)),material_kind=dict(collections.Counter(r['material_kind'] for r in rows)),geography=dict(collections.Counter('国内' if r['country']=='中国' else ('国际体系' if not r['country'] else '国外国家') for r in rows)),layout_and_full_text_invariance_claimed=False,core_values_identical_within_groups=True,new_release_images_generated=0,new_release_images_accepted=0,new_prompt_individually_approved=0,category_research_repeated=False)

def reports(repo,bases,rows):
    dump_json(repo/'reports/qa_report.json',stats(bases,rows))
    text_file(repo/'reports/qa_report.md','# 20261021_v2执行统计\n\n1430类、4290条，实体4254条、电子36条。每组只保证给定核心取值一致；允许版式、固定印刷正文与装饰不同。此处报告的是提示词和导出检查，不是图像验收。\n\n本轮未运行生图。用户此前小样的反馈范围和限制见docs/pilot_feedback.md；不将小样认可转成4290条正文或图像通过。\n')
    dump_json(repo/'reports/core_value_consistency.json',dict(policy='A',groups=len(bases),same_core_fields=True,same_target_texts=True,full_visible_text_invariance_required=False,layout_invariance_required=False,per_group=[dict(category_entry_id=b['category_entry_id'],core_value_group_id=b['core_value_group_id'],core_values_sha256=b['core_values_sha256']) for b in bases]))
    dump_jl(repo/'reports/review_status.jsonl',[dict(sample_id=r['sample_id'],prompt_sha256=r['prompt_sha256'],policy_A_selected_by_user=True,new_prompt_individually_approved=False,new_release_image_generated=False,new_release_image_accepted=False) for r in rows])
    dump_jl(repo/'provenance/parent_mapping.jsonl',[{k:r[k] for k in ['sample_id','category_entry_id','source_base_sample_id','source_prompt_sha256','source_release_version','core_value_group_id','core_values_sha256']} for r in rows])
    ids=['CAT-B0198','CAT-A0030','CAT-D0081','CAT-C0207']
    text_file(repo/'docs/examples_12.md','# 四类文书的12条A口径提示词\n\n仅为文字导出，本版本未出图。同组给定值相同，版式和固定文案允许变化。\n\n'+'\n\n'.join(f'## {r["category_name"]} / {r["capture_label"]}\n\n{r["sample_id"]}\n\n```text\n{r["prompt"]}\n```' for id in ids for r in rows if r['category_entry_id']==id)+'\n')

def build(repo):
    bases,rows=generate(repo);export(repo,rows);reports(repo,bases,rows);return stats(bases,rows)

def validate(repo,report_dir=None,rebuild=False):
    bases=read_jl(repo/'data/base_values_1430.jsonl');rows=read_jl(repo/'data/prompts.jsonl');taxonomy=read_json(repo/'config/taxonomy.json');policy=read_json(repo/'config/prompt_policy.json');config=read_json(repo/'config/imaging_rules.json')
    groups=collections.defaultdict(list)
    for r in rows:groups[r['core_value_group_id']].append(r)
    by={b['core_value_group_id']:b for b in bases};ck={}
    ck['1430 unique categories and exact taxonomy']=len(bases)==len(by)==1430 and {r['category_entry_id'] for r in rows}=={x['entry_id'] for x in taxonomy['entries']}
    ck['4290 unique IDs in versioned range']=[r['sample_id'] for r in rows]==[f'CRD-20261021-v2-{i:05d}' for i in range(1,4291)]
    ck['5 first level and 40 second level']=len({r['category_l1_id'] for r in rows})==5 and len({r['category_l2_id'] for r in rows})==40
    ck['three rows per supplied-value group']=set(groups)==set(by) and all(len(rs)==3 for rs in groups.values())
    ck['physical and electronic mode sets']=all({r['capture_mode'] for r in rs}==set(ELECTRONIC if by[k]['material_kind']=='electronic' else PHYSICAL) for k,rs in groups.items())
    ck['4254 physical and 36 electronic']=sum(r['material_kind']=='electronic' for r in rows)==36 and len(rows)==4290
    ck['core payload hashes valid']=all(sha(canonical(core_payload(b)).encode())==b['core_values_sha256'] for b in bases)
    ck['same supplied fields targets and digest']=all(all(r['core_fields']==b['core_fields'] and r['target_texts']==b['target_texts'] and r['expected_core_values']==b['expected_core_values'] and r['core_values_sha256']==b['core_values_sha256'] for r in groups[k]) for k,b in by.items())
    ck['knowledge-guided base regenerated']=all(b['base_prompt']==base_prompt(b,policy) for b in bases)
    ck['only capture paragraph varies within group']=all(r['prompt']==render(by[r['core_value_group_id']],r['capture_mode'],config) for r in rows)
    ck['all supplied values embedded']=all(all(f['value'] and f['value'] in r['prompt'] for f in r['core_fields']) for r in rows)
    ck['current prompt SHA256']=all(sha(r['prompt'].encode())==r['prompt_sha256'] for r in rows)
    ck['A semantics explicit no full-page lock']=all(r['group_semantics']==SEMANTICS and not r['layout_locked'] and not r['full_visible_text_locked'] and r['target_texts_scope']=='supplied_core_values_only' for r in rows)
    ck['no inherited layout skeleton or forced document ratio']=all('【文书骨架】' not in r['prompt'] and '表格每行依次为' not in r['prompt'] and '文档本体宽:高=' not in r['prompt'] and r['document_aspect_ratio'] is None for r in rows)
    ck['all clean with original team group semantics']=all(r['variant']=='clean' and not r['group_id'] and not r['tamper_types'] and not r['tamper_details'] for r in rows)
    ck['standard canvas agrees with prompt']=all(r['aspect_ratio'] in {'3:2','2:3','1:1'} and f'输出画布宽:高={r["aspect_ratio"]}' in r['prompt'] for r in rows)
    caps={r['sample_id']:r['prompt'].split('【本条成像要求】')[1] for r in rows}
    ck['electronic never asks for a scanner']=all('平板扫描仪' not in r['prompt'] for r in rows if r['material_kind']=='electronic')
    ck['camera modes do not inherit scanner restrictions']=all('平板扫描仪' not in caps[r['sample_id']] for r in rows if r['capture_mode']!='scan')
    ck['no new-release images claimed']=all(not r['new_release_image_generated'] and not r['new_release_image_accepted'] and not r['new_prompt_individually_approved'] for r in rows)
    cs=read_csv(repo/'data/credential_prompts_4290_unified.csv');ck['CSV exact values and metadata']=cs==team_rows(rows)
    expected='\n\n'.join(f'===== {r["sample_id"]} | {r["category_name"]} | {r["capture_label"]} =====\n{r["prompt"]}' for r in rows)+'\n';ck['TXT exact prompts']=(repo/'data/prompts.txt').read_text(encoding='utf-8')==expected
    from openpyxl import load_workbook
    w=load_workbook(repo/'data/prompts_filter.xlsx',read_only=True,data_only=True);cells=list(w.active.values);w.close();keys=list(cells[0]);ck['Excel exact rows']=len(cells)==4291 and all([v if v is not None else '' for v in vals]==[r[k] for k in keys] for vals,r in zip(cells[1:],rows))
    ck['statistics computed from rows']=read_json(repo/'reports/qa_report.json')==stats(bases,rows)
    ck['12 sample prompts present']=(repo/'docs/examples_12.md').read_text(encoding='utf-8').count('```text')==12
    if rebuild:
        _,rebuilt=generate(repo);ck['offline JSONL byte reproduction']=(''.join(canonical(r)+'\n' for r in rebuilt)).encode()==(repo/'data/prompts.jsonl').read_bytes()
        buf=io.StringIO(newline='');w=csv.DictWriter(buf,COLUMNS);w.writeheader();w.writerows(team_rows(rebuilt));ck['offline CSV byte reproduction']=b'\xef\xbb\xbf'+buf.getvalue().encode()==(repo/'data/credential_prompts_4290_unified.csv').read_bytes()
    result=dict(passed=all(ck.values()),checks=ck,scope='prompt data and exports only; no category re-audit or image acceptance')
    if report_dir:dump_json(pathlib.Path(report_dir)/'validation.json',result)
    if not result['passed']:raise AssertionError([k for k,v in ck.items() if not v])
    return result

def seal(repo):
    files=[p for p in sorted(repo.rglob('*')) if p.is_file() and p.name not in {'release_manifest.json','SHA256SUMS.txt'} and '__pycache__' not in p.parts]
    payload=[dict(path=p.relative_to(repo).as_posix(),bytes=p.stat().st_size,sha256=sha(p.read_bytes())) for p in files]
    dump_json(repo/'release_manifest.json',dict(version=VERSION,payload_file_count=len(payload),payload_bytes_total=sum(p['bytes'] for p in payload),payload_files=payload))
    text_file(repo/'SHA256SUMS.txt',''.join(p['sha256']+'  '+p['path']+'\n' for p in payload));return len(files)

def verify_hashes(repo):
    m=read_json(repo/'release_manifest.json');expected={r['path']:r for r in m['payload_files']};files={p.relative_to(repo).as_posix():p for p in repo.rglob('*') if p.is_file() and p.name not in {'release_manifest.json','SHA256SUMS.txt'} and '__pycache__' not in p.parts}
    assert set(expected)==set(files),'manifest coverage'
    assert all(sha(files[k].read_bytes())==v['sha256'] and files[k].stat().st_size==v['bytes'] for k,v in expected.items()),'payload hashes'
    assert m['payload_file_count']==len(files) and m['payload_bytes_total']==sum(p.stat().st_size for p in files.values())
    assert (repo/'SHA256SUMS.txt').read_text(encoding='utf-8')==''.join(r['sha256']+'  '+r['path']+'\n' for r in m['payload_files'])
    return dict(passed=True,files=len(files))
