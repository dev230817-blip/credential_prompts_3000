"""Build, export and validate the versioned 4290-prompt release, offline."""
from __future__ import annotations
import collections,copy,csv,datetime,hashlib,io,json,pathlib,re,zipfile

VERSION='credential-category-imaging-20261021-v1'
COLUMNS=['sample_id','direction','category_l1','category_l2','variant','aspect_ratio','prompt','group_id','target_texts','tamper_types','tamper_details','extra_json']
PHYSICAL=('scan','front_photo','oblique_photo')
ELECTRONIC=('digital_flat','screen_front_photo','screen_oblique_photo')

def sha(data):return hashlib.sha256(data).hexdigest()
def canonical(value):return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':'))
def read_json(p):return json.loads(p.read_text(encoding='utf-8'))
def read_jl(p):return [json.loads(x) for x in p.read_text(encoding='utf-8').splitlines() if x.strip()]
def dump_json(p,value):
    p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def dump_jl(p,rows):
    p.parent.mkdir(parents=True,exist_ok=True);p.write_text(''.join(canonical(x)+'\n' for x in rows),encoding='utf-8')
def dump_csv(p,rows,fields=None):
    p.parent.mkdir(parents=True,exist_ok=True)
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fields or list(rows[0]));writer.writeheader();writer.writerows(rows)
def read_csv(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))

def content_payload(base):
    return {k:base[k] for k in ['category_entry_id','fields','expected_visible_text','target_texts','fixed_sections','document_aspect_ratio','material_kind','seal_policy']}

def render(base,mode,rules):
    paragraph=rules['modes'][mode]['text'].format(document_aspect_ratio=base['document_aspect_ratio'],canvas_aspect_ratio=base['canvas_aspect_ratio'],carrier=base['carrier_label'])
    sections=base['fixed_sections'][:]
    sections.insert(1,'【图像与载体】\n'+paragraph)
    return '\n\n'.join(sections)

def generate(repo):
    bases=read_jl(repo/'data/base_content_1430.jsonl');rules=read_json(repo/'config/imaging_rules.json')
    rows=[]
    for base in bases:
        for mode in ELECTRONIC if base['material_kind']=='electronic' else PHYSICAL:
            ix=len(rows)+1;prompt=render(base,mode,rules)
            r={k:copy.deepcopy(base[k]) for k in ['category_entry_id','category_l1_id','category_l1','category_l2_id','category_l2','category_name','country','region','document_family','family_label','material_kind','seal_policy','instruction_language','visible_language','language_basis','fields','expected_visible_text','target_texts','source_base_sample_id','source_prompt_sha256','source_version','content_group_id','content_sha256','document_aspect_ratio','canvas_aspect_ratio','source_ids','evidence_status','layout_status']}
            r.update(sample_id=f'CRD-20261021-v1-{ix:05d}',version=VERSION,release_version='20261021_v1',direction='credential',variant='clean',group_id='',tamper_types='',tamper_details='',capture_mode=mode,capture_label=rules['modes'][mode]['label'],aspect_ratio=base['canvas_aspect_ratio'],aspect_ratio_basis='标准生成画布；文档本体比例另存，不拉伸载体',prompt=prompt,prompt_sha256=sha(prompt.encode('utf-8')),image_generated=False,image_accepted=False,manual_body_approved=False,review_status='approved_scope_and_family_feedback; new_capture_bodies_not_individually_approved',layout_basis='inherited_synthetic_family_design_not_official_layout_confirmation')
            rows.append(r)
    return bases,rows

def team_rows(rows):
    result=[]
    for r in rows:
        x={k:r[k] for k in COLUMNS if k!='extra_json'}
        extra={k:v for k,v in r.items() if k not in set(COLUMNS)|{'fields'}}
        extra['fields']=r['fields']
        x['extra_json']=canonical(extra);result.append(x)
    return result

def deterministic_xlsx(p,rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    wb=Workbook();ws=wb.active;ws.title='提示词4290'
    keys=['sample_id','category_entry_id','category_l1','category_l2','category_name','country','region','material_kind','capture_mode','aspect_ratio','document_aspect_ratio','content_group_id','prompt']
    ws.append(keys)
    for r in rows:ws.append([r[k] for k in keys])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='274D72')
    for col,width in zip('ABCDEFGHIJKLM',[31,19,20,33,60,16,18,15,24,14,20,36,110]):ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):
        row[-1].alignment=Alignment(wrap_text=True,vertical='top')
    wb.properties.creator='credential-prompt-project';wb.properties.lastModifiedBy='credential-prompt-project'
    fixed=datetime.datetime(2026,10,21,0,0,0)
    wb.properties.created=fixed;wb.properties.modified=fixed
    memory=io.BytesIO();wb.save(memory)
    # openpyxl writes current modified time; canonicalize archive metadata and core time.
    with zipfile.ZipFile(io.BytesIO(memory.getvalue())) as src,zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED) as out:
        for name in sorted(src.namelist()):
            data=src.read(name)
            if name=='docProps/core.xml':
                data=re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)',rb'\g<1>2026-10-21T00:00:00Z\2',data)
            info=zipfile.ZipInfo(name,(2026,10,21,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;info.external_attr=0o600<<16
            out.writestr(info,data)

def export(repo,rows):
    data=repo/'data';data.mkdir(exist_ok=True)
    dump_jl(data/'prompts.jsonl',rows)
    dump_csv(data/'credential_prompts_4290_unified.csv',team_rows(rows),COLUMNS)
    (data/'prompts.txt').write_text('\n\n'.join(f'===== {r["sample_id"]} | {r["category_entry_id"]} | {r["capture_label"]} | {r["category_name"]} =====\n{r["prompt"]}' for r in rows)+'\n',encoding='utf-8')
    deterministic_xlsx(data/'prompts_filter.xlsx',rows)

def stats_for(bases,rows):
    return dict(version=VERSION,base_content_count=len(bases),category_count=len({x['category_entry_id'] for x in bases}),prompt_count=len(rows),l1=dict(collections.Counter(x['category_l1'] for x in rows)),capture_modes=dict(collections.Counter(x['capture_mode'] for x in rows)),geography=dict(collections.Counter('国内' if x['country']=='中国' else ('国际体系' if not x['country'] else '国外国家') for x in rows)),material_kind=dict(collections.Counter(x['material_kind'] for x in rows)),all_clean=True,new_image_generated=0,new_image_accepted=0,manual_body_approved_count=0,classification_review_repeated=False,source_research_performed=False)

def write_reports(repo,bases,rows):
    reports=repo/'reports';reports.mkdir(exist_ok=True);stats=stats_for(bases,rows)
    dump_json(reports/'qa_report.json',stats)
    text=['# 4290条数据统计与边界','','每个具体类别使用一份完整基础内容、三种成像；类别数不等于已核视觉模板数。','',f'基础内容{len(bases)}份，提示词{len(rows)}条；实体4254、电子36。','', '| 成像 | 数量 |','| --- | ---: |']
    text += [f'| {k} | {v} |' for k,v in stats['capture_modes'].items()]
    text += ['','同组三条fields、expected_visible_text、target_texts和content_sha256一致。source_prompt_sha256绑定父版本提示词。','', '类别和已有知识支持范围沿用；未重新检索、未重审类别。所有文档比例、载体和版式沿用已定合成设计，不冒充官方精确规格。','', '新正文逐条人工确认0；图片生成0、图片验收0。原六图反馈只支持此前调整方向，不作为4290条图片已验收。','']
    (reports/'qa_report.md').write_text('\n'.join(text),encoding='utf-8')
    dump_json(reports/'content_invariance.json',dict(groups=len(bases),rows=len(rows),same_fields=True,same_target_texts=True,same_content_sha256=True,content_hash_basis='canonical content payload; excludes capture paragraph',per_group=[dict(content_group_id=b['content_group_id'],category_entry_id=b['category_entry_id'],content_sha256=b['content_sha256'],samples=[r['sample_id'] for r in rows if r['content_group_id']==b['content_group_id']]) for b in bases]))
    dump_jl(reports/'review_status.jsonl',[dict(sample_id=r['sample_id'],prompt_sha256=r['prompt_sha256'],category_scope_accepted=True,new_capture_body_individually_approved=False,image_generated=False,image_accepted=False) for r in rows])
    examples=['CAT-B0013','CAT-B0014','CAT-B0198','CAT-A0030','CAT-D0009','CAT-A0264']
    sample_rows=[r for k in examples for r in rows if r['category_entry_id']==k]
    (repo/'docs/examples_18.md').write_text('# 六种文书的18条完整提示词\n\n仅为文字样例，未运行本轮生图。\n\n'+'\n\n'.join(f'## {r["category_name"]} / {r["capture_label"]}\n\n{r["sample_id"]}\n\n```text\n{r["prompt"]}\n```' for r in sample_rows)+'\n',encoding='utf-8')
    dump_jl(repo/'provenance/parent_mapping.jsonl',[dict(sample_id=r['sample_id'],category_entry_id=r['category_entry_id'],source_base_sample_id=r['source_base_sample_id'],source_version=r['source_version'],source_prompt_sha256=r['source_prompt_sha256'],content_group_id=r['content_group_id'],content_sha256=r['content_sha256']) for r in rows])

def build(repo):
    bases,rows=generate(repo);export(repo,rows);write_reports(repo,bases,rows)
    return stats_for(bases,rows)

def validate(repo,report_dir=None,rebuild=False):
    bases=read_jl(repo/'data/base_content_1430.jsonl');rows=read_jl(repo/'data/prompts.jsonl');taxonomy=read_json(repo/'config/taxonomy.json');rules=read_json(repo/'config/imaging_rules.json')
    groups=collections.defaultdict(list)
    for r in rows:groups[r['content_group_id']].append(r)
    base_by={b['content_group_id']:b for b in bases}
    ck={}
    ck['4290 unique sample IDs']=len(rows)==len({r['sample_id'] for r in rows})==4290
    ck['1430 bases and taxonomy coverage']=len(bases)==len(base_by)==1430 and {r['category_entry_id'] for r in rows}=={e['entry_id'] for e in taxonomy['entries']}
    ck['5 first-level and 40 second-level']=len({r['category_l1_id'] for r in rows})==5 and len({r['category_l2_id'] for r in rows})==40
    ck['three rows per content group']=set(groups)==set(base_by) and all(len(xs)==3 for xs in groups.values())
    ck['physical and electronic modes']=all({r['capture_mode'] for r in xs}==set(ELECTRONIC if base_by[k]['material_kind']=='electronic' else PHYSICAL) for k,xs in groups.items())
    ck['4254 physical and 36 electronic']=sum(r['material_kind']=='electronic' for r in rows)==36 and sum(b['material_kind']=='electronic' for b in bases)==12
    ck['content digest matches frozen payload']=all(sha(canonical(content_payload(b)).encode())==b['content_sha256'] for b in bases)
    ck['same fields and visible values in group']=all(all(r['fields']==b['fields'] and r['expected_visible_text']==b['expected_visible_text'] and r['target_texts']==b['target_texts'] and r['content_sha256']==b['content_sha256'] for r in groups[k]) for k,b in base_by.items())
    ck['only capture paragraph changes']=all(r['prompt']==render(base_by[r['content_group_id']],r['capture_mode'],rules) for r in rows)
    ck['current prompt hashes']=all(sha(r['prompt'].encode())==r['prompt_sha256'] for r in rows)
    ck['complete field values present']=all(all(x['value'] and x['value'] in r['prompt'] for x in r['fields']) for r in rows)
    ck['clean variants, no tamper groups']=all(r['variant']=='clean' and not r['group_id'] and not r['tamper_types'] and not r['tamper_details'] for r in rows)
    ck['standard canvases, separate document proportions']=all(r['aspect_ratio'] in {'3:2','2:3','1:1'} and f'输出画布宽:高={r["aspect_ratio"]}' in r['prompt'] and r['document_aspect_ratio']==base_by[r['content_group_id']]['document_aspect_ratio'] for r in rows)
    caps={r['sample_id']:r['prompt'].split('【图像与载体】')[1].split('【文书骨架】')[0] for r in rows}
    ck['electronic uses screen or native page, no scanner']=all('平板扫描仪' not in caps[r['sample_id']] and '载体为电子页面' in caps[r['sample_id']] for r in rows if r['material_kind']=='electronic')
    ck['photos do not inherit flatbed restrictions']=all('平板扫描仪' not in caps[r['sample_id']] and '没有拍摄透视' not in caps[r['sample_id']] for r in rows if r['capture_mode']!='scan')
    ck['new image state remains zero']=all(r['image_generated'] is False and r['image_accepted'] is False and r['manual_body_approved'] is False for r in rows)
    cs=read_csv(repo/'data/credential_prompts_4290_unified.csv');ck['CSV exact content and metadata']=cs==team_rows(rows)
    expected='\n\n'.join(f'===== {r["sample_id"]} | {r["category_entry_id"]} | {r["capture_label"]} | {r["category_name"]} =====\n{r["prompt"]}' for r in rows)+'\n'
    ck['TXT exact content']=(repo/'data/prompts.txt').read_text(encoding='utf-8')==expected
    from openpyxl import load_workbook
    wb=load_workbook(repo/'data/prompts_filter.xlsx',read_only=True,data_only=True);cells=list(wb.active.values);wb.close()
    keys=list(cells[0]);ck['Excel exact row values']=len(cells)==4291 and all([v if v is not None else '' for v in vals]==[r[k] for k in keys] for vals,r in zip(cells[1:],rows))
    qa=read_json(repo/'reports/qa_report.json');ck['reported counts computed from data']=qa==stats_for(bases,rows)
    if rebuild:
        _,rebuilt=generate(repo);ck['offline generation reproduces JSONL']=rebuilt==rows
        buf=io.StringIO(newline='');w=csv.DictWriter(buf,COLUMNS);w.writeheader();w.writerows(team_rows(rebuilt))
        ck['CSV bytes reproducible']=b'\xef\xbb\xbf'+buf.getvalue().encode()==(repo/'data/credential_prompts_4290_unified.csv').read_bytes()
    result=dict(passed=all(ck.values()),checks=ck,scope='generation, content invariance, exports; no category or source re-audit')
    if report_dir:dump_json(pathlib.Path(report_dir)/'validation.json',result)
    if not result['passed']:raise AssertionError([k for k,v in ck.items() if not v])
    return result

def seal(repo):
    excluded={'release_manifest.json','SHA256SUMS.txt'}
    files=[p for p in sorted(repo.rglob('*')) if p.is_file() and p.name not in excluded and '__pycache__' not in p.parts]
    payload=[dict(path=p.relative_to(repo).as_posix(),bytes=p.stat().st_size,sha256=sha(p.read_bytes())) for p in files]
    dump_json(repo/'release_manifest.json',dict(version=VERSION,payload_file_count=len(payload),payload_bytes_total=sum(x['bytes'] for x in payload),payload_files=payload))
    (repo/'SHA256SUMS.txt').write_text(''.join(x['sha256']+'  '+x['path']+'\n' for x in payload),encoding='utf-8')
    return len(payload)

def verify_hashes(repo):
    m=read_json(repo/'release_manifest.json');expected={x['path']:x for x in m['payload_files']}
    actual={p.relative_to(repo).as_posix():p for p in repo.rglob('*') if p.is_file() and p.name not in {'release_manifest.json','SHA256SUMS.txt'} and '__pycache__' not in p.parts}
    assert set(actual)==set(expected),'manifest coverage'
    assert all(sha(actual[k].read_bytes())==v['sha256'] and actual[k].stat().st_size==v['bytes'] for k,v in expected.items()),'manifest bytes'
    assert len(expected)==m['payload_file_count'] and sum(v['bytes'] for v in expected.values())==m['payload_bytes_total']
    assert (repo/'SHA256SUMS.txt').read_text(encoding='utf-8')==''.join(x['sha256']+'  '+x['path']+'\n' for x in m['payload_files'])
    return dict(passed=True,files=len(expected))
