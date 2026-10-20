"""Re-export the release JSONL to CSV, TXT and a deterministic Excel workbook."""
import argparse
import csv
from datetime import datetime
import io
import json
import os
from pathlib import Path
import re
import zipfile

# Use the same XML writer whether optional lxml happens to be installed or not.
# Set this before openpyxl is imported by this CLI or its validator.
os.environ['OPENPYXL_LXML']='False'

COLUMNS=['sample_id','direction','category_l1','category_l2','variant','aspect_ratio','prompt','group_id','target_texts','tamper_types','tamper_details','extra_json']

def export(repo,out):
    repo=Path(repo).resolve();out=Path(out).resolve()
    if out==repo or out.is_relative_to(repo):raise ValueError('Output must be outside the release directory')
    out.mkdir(parents=True,exist_ok=False)
    rows=[json.loads(s) for s in (repo/'data/prompts.jsonl').read_text(encoding='utf-8').splitlines() if s.strip()]
    with (out/'credential_prompts_3000_unified.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=COLUMNS);w.writeheader()
        for row in rows:
            record={k:row[k] for k in COLUMNS}
            record['extra_json']=json.dumps(row['extra_json'],ensure_ascii=False,separators=(',',':'))
            w.writerow(record)
    (out/'prompts.txt').write_text(''.join(f"[{r['sample_id']}]\n{r['prompt']}\n\n" for r in rows),encoding='utf-8',newline='\n')
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb=Workbook();ws=wb.active;ws.title='提示词'
    ws.append(['编号','一级类别','二级类型','国家','成像方式','指令语言','图中文字语言','对象ID','模板ID','原组ID','画幅','完整提示词','正文人工状态','图片状态'])
    for r in rows:
        m=r['extra_json'];ws.append([r['sample_id'],r['category_l1'],r['category_l2'],m['country'],m['capture_label'],m['instruction_language'],m['visible_language'],m['entity_id'],m['profile_id'],m['source_group_id'],r['aspect_ratio'],r['prompt'],m['current_body_review_status'],'not_generated_not_accepted'])
    taxonomy=json.loads((repo/'config/taxonomy.json').read_text(encoding='utf-8'))
    ws=wb.create_sheet('对象分类');ws.append(['对象ID','名称','一级','二级','版本','页面'])
    for c in taxonomy['categories']:
        for sub in c['subcategories']:
            for e in sub['current_examples']:ws.append([e['entity_id'],e['name_zh'],c['name_zh'],sub['name_zh'],e['selected_edition'],e['surface']])
    ws=wb.create_sheet('成像模板');ws.append(['模板ID','对象ID','成像方式','画幅','对象描述','载体描述','版式描述','渲染描述'])
    templates=json.loads((repo/'config/imaging_templates_162.json').read_text(encoding='utf-8'))['templates']
    for t in templates:
        s=t['shared_prompt_sections'];ws.append([t['profile_id'],t['entity_id'],t['capture_mode'],t['aspect_ratio'],s['object'],s['carrier'],s['layout'],s['rendering']])
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=22
        for row in ws:
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
        ws.row_dimensions[1].height=30
    wb['提示词'].column_dimensions['L'].width=100
    wb.properties.creator='credential_prompts_3000';wb.properties.lastModifiedBy='credential_prompts_3000'
    wb.properties.created=datetime(2026,10,20);wb.properties.modified=datetime(2026,10,20)
    buf=io.BytesIO();wb.save(buf);wb.close()
    # Normalise ZIP timestamps and the OOXML modified property for reproducible bytes.
    with zipfile.ZipFile(buf) as src,zipfile.ZipFile(out/'prompts_filter.xlsx','w',compression=zipfile.ZIP_STORED) as dest:
        for name in sorted(src.namelist()):
            data=src.read(name)
            if name=='docProps/core.xml':data=re.sub(rb'(<(?:[A-Za-z0-9_]+:)?modified\b[^>]*>).*?(</(?:[A-Za-z0-9_]+:)?modified>)',rb'\g<1>2026-10-20T00:00:00Z\g<2>',data)
            # Stored members avoid zlib-version differences between platforms.
            info=zipfile.ZipInfo(name,(2026,10,20,0,0,0));info.compress_type=zipfile.ZIP_STORED
            info.create_system=0;info.external_attr=0o600 << 16
            dest.writestr(info,data)
    return {'rows':len(rows),'files':sorted(p.name for p in out.iterdir())}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--repo',required=True);ap.add_argument('--out',required=True);a=ap.parse_args()
    print(json.dumps(export(a.repo,a.out),ensure_ascii=False))

if __name__=='__main__':main()
