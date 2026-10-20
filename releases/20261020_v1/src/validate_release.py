"""Validate released data, cross-file bindings, exports and declared review scope."""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import tempfile

from build_qa_report import make_reports
from export_data import export, COLUMNS
from verify_hashes import verify, files

VERSION='credential-imaging-20261020-v1'

def j(path):return json.loads(Path(path).read_text(encoding='utf-8-sig'))
def jl(path):return [json.loads(s) for s in Path(path).read_text(encoding='utf-8-sig').splitlines() if s.strip()]
def h(text):return hashlib.sha256(text.encode('utf-8')).hexdigest()
def require(value,message='Condition failed'):
    if not value:raise ValueError(message)

def validate(repo,check_hashes=True):
    repo=Path(repo).resolve();results=[]
    def check(code,label,fn):
        try:fn();results.append({'id':code,'label':label,'passed':True})
        except Exception as exc:results.append({'id':code,'label':label,'passed':False,'error':str(exc)})
    try:
        rows=jl(repo/'data/prompts.jsonl');bases=jl(repo/'data/clean_content_1000.jsonl')
        taxonomy=j(repo/'config/taxonomy.json');keywords=j(repo/'config/keywords.json')
        ts=j(repo/'config/imaging_templates_162.json')['templates'];rules=j(repo/'config/scan_rules.json')
        changes=j(repo/'reports/scan_changes_980.json');sources=j(repo/'config/source_index.json')
        reviews=jl(repo/'reports/review_status.jsonl');overrides=j(repo/'provenance/approved_parent_overrides.json')
    except Exception as exc:return {'passed':False,'checks':0,'results':[{'id':'LOAD','passed':False,'error':str(exc)}]}
    bybase={r['source_group_id']:r for r in bases}
    bytemplate={(t['profile_id'],t['capture_mode']):t for t in ts}
    change_map={r['sample_id']:r for r in changes['rows']}
    check('V01','Exactly 3,000 sequential unique CRD IDs',lambda:require([r['sample_id'] for r in rows]==[f'CRD-{i:04}' for i in range(1,3001)]))
    check('V02','1,000 unique base groups, three capture modes each',lambda:require(len(bases)==len(bybase)==1000 and Counter(r['extra_json']['source_group_id'] for r in rows)==Counter({g:3 for g in bybase}) and all({r['extra_json']['capture_mode'] for r in rows if r['extra_json']['source_group_id']==g}=={'scan','front_photo','oblique_photo'} for g in bybase)))
    check('V03','Clean-only team schema and empty tamper/group fields',lambda:require(all(r['direction']=='credential' and r['variant']=='clean' and r['group_id']==r['tamper_types']==r['tamper_details']=='' and all(k in r for k in COLUMNS) for r in rows)))
    check('V04','50 entities, 54 profiles, 162 unique capture templates',lambda:require(len({r['extra_json']['entity_id'] for r in rows})==50 and len({r['extra_json']['profile_id'] for r in rows})==54 and len(ts)==len(bytemplate)==162))
    check('V05','Physical/electronic capture distribution',lambda:require(Counter((r['extra_json']['material_kind']=='electronic',r['extra_json']['capture_mode']) for r in rows)==Counter({(electronic,mode):(20 if electronic else 980) for electronic in [True,False] for mode in ['scan','front_photo','oblique_photo']})))
    check('V06','All current prompt hashes and version fields',lambda:require(all(r['prompt_sha256']==h(r['prompt']) and r['version']==VERSION for r in rows) and len({r['prompt'] for r in rows})==3000))
    check('V07','Historical base source-body hashes',lambda:require(all(b['source_prompt_sha256']==h(b['source_prompt']) for b in bases)))
    def check_inheritance():
        for r in rows:
            m=r['extra_json'];b=bybase[m['source_group_id']]
            require(r['expected_visible_text']==b['expected_visible_text'] and r['target_texts']=='|'.join(str(v) for v in b['expected_visible_text'].values()),r['sample_id']+' visible values')
            require(all(m[k]==b[k] for k in ['entity_id','profile_id','source_prompt_id','source_prompt_sha256','country','surface','material_kind','source_ids','instruction_language','visible_language','pending_checks']),r['sample_id']+' metadata')
            require(r['aspect_ratio']==b['aspect_ratio'] and r['category_l1']==b['category_l1'] and r['category_l2']==b['category_l2'],r['sample_id']+' classification/aspect')
    check('V08','Every current row inherits its base values and provenance',check_inheritance)
    def check_overrides():
        require(len(overrides['changes'])==3 and sum(len(v) for v in overrides['changes'].values())==6)
        for g,b in bybase.items():
            expected=dict(b['source_expected_visible_text'])
            for field,c in overrides['changes'].get(g,{}).items():
                require(expected[field]==c['old']);expected[field]=c['new']
            require(expected==b['expected_visible_text'],g)
    check('V09','Only the six approved base-value replacements',check_overrides)
    def check_template_binding():
        for r in rows:
            m=r['extra_json'];t=bytemplate[(m['profile_id'],m['capture_mode'])]
            require(t['entity_id']==m['entity_id'] and t['aspect_ratio']==r['aspect_ratio'] and t['template_version']==VERSION,r['sample_id'])
            require(all(section in r['prompt'] for section in t['shared_prompt_sections'].values()),r['sample_id']+' template sections')
    check('V10','All current prompt sections match the imaging templates',check_template_binding)
    check('V11','Scan wording restricted to physical scans',lambda:require(all((rules['zh' if r['extra_json']['instruction_language']=='中文' else 'en'] in r['prompt'])==(r['extra_json']['capture_mode']=='scan' and r['extra_json']['material_kind']!='electronic') for r in rows)))
    def check_changes():
        physical=[r for r in rows if r['extra_json']['capture_mode']=='scan' and r['extra_json']['material_kind']!='electronic']
        require(changes['version']==VERSION and changes['baseline']=='20261014_v3' and changes['changed_rows']==len(changes['rows'])==len(change_map)==len(physical)==980)
        require(set(change_map)=={r['sample_id'] for r in physical})
        require(dict(Counter(r['extra_json']['instruction_language'] for r in physical))==changes['language_counts'])
        for r in physical:
            m=r['extra_json'];c=change_map[r['sample_id']];rev=m['capture_revision']
            require(c['new_prompt_sha256']==r['prompt_sha256'] and c['old_prompt_sha256']==rev['baseline_prompt_sha256'] and c['old_prompt_sha256']!=c['new_prompt_sha256'])
            require(all(c[k]==m[k] for k in ['entity_id','profile_id','instruction_language','material_kind']))
            require(rev['scope']=='physical_scan_capture_only' and rev['version']==VERSION)
    check('V12','980 change records bind current IDs and recorded v3 hashes',check_changes)
    def check_taxonomy():
        categories=taxonomy['categories'];entries=[(c,s,e) for c in categories for s in c['subcategories'] for e in s['current_examples']]
        mapping={e['entity_id']:(c['name_zh'],s['name_zh']) for c,s,e in entries}
        require(len(categories)==5 and sum(len(c['subcategories']) for c in categories)==29 and len(entries)==len(mapping)==50)
        require(all(mapping[r['extra_json']['entity_id']]==(r['category_l1'],r['category_l2']) for r in rows))
    check('V13','Taxonomy maps every entity exactly once',check_taxonomy)
    check('V14','Keyword options bind the released profiles',lambda:require(keywords['version']==VERSION and len(keywords['profiles'])==54 and {p['profile_id'] for p in keywords['profiles']}=={t['profile_id'] for t in ts} and all(p['allowed_modes']==['scan','front_photo','oblique_photo'] for p in keywords['profiles'])))
    def check_reviews():
        require(len(reviews)==3000 and len({r['sample_id'] for r in reviews})==3000)
        for r,v in zip(rows,reviews):
            require(v['sample_id']==r['sample_id'] and v['prompt_sha256']==r['prompt_sha256'])
            require(r['manual_body_approved'] is False and v['manual_body_approved'] is False)
            require(v['current_body_review_status']==r['extra_json']['current_body_review_status']=='pending_user_spot_review')
            require(v['source_body_approved']==r['extra_json']['source_body_approved'])
    check('V15','Current review records are hash-bound and remain pending',check_reviews)
    check('V16','No generated or accepted images asserted',lambda:require(all(r['image_generated'] is False and r['image_accepted'] is False and r['extra_json']['current_image_generated'] is False and r['extra_json']['current_image_accepted'] is False for r in rows) and all(v['image_generated'] is False and v['image_accepted'] is False for v in reviews)))
    def check_sources():
        used={s for r in rows for s in r['extra_json']['source_ids']};ids=[s['source_id'] for s in sources['sources']]
        require(len(ids)==len(set(ids)) and set(ids)==used and sources['unresolved_source_ids']==[])
        require('not rechecked' in sources['scope'])
    check('V17','All referenced source IDs resolve to inherited public metadata',check_sources)
    def check_exports():
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory(prefix='credential_export_') as tmp:
            dest=Path(tmp)/'out';export(repo,dest)
            for name in ['credential_prompts_3000_unified.csv','prompts.txt','prompts_filter.xlsx']:
                require((dest/name).read_bytes()==(repo/'data'/name).read_bytes(),'Export mismatch: '+name)
        wb=load_workbook(repo/'data/prompts_filter.xlsx',read_only=True,data_only=False)
        require(wb.sheetnames==['提示词','对象分类','成像模板'])
        require([wb[s].max_row for s in wb.sheetnames]==[3001,51,163]);wb.close()
    check('V18','CSV/TXT/Excel are deterministic same-version exports',check_exports)
    def check_qa():
        q,text=make_reports(repo)
        require(j(repo/'reports/qa_report.json')==q,'QA JSON differs from data')
        require((repo/'reports/qa_report.md').read_text(encoding='utf-8')==text,'QA Markdown differs from data')
    check('V19','QA reports reconstruct from public data',check_qa)
    def check_privacy():
        pattern=re.compile(r'(?<![A-Za-z0-9])[A-Za-z]:[\\/]|file://|/(?:Users|home)/|github_pat_[A-Za-z0-9_]{20,}|gh[pousr]_[A-Za-z0-9]{20,}')
        def strings(obj):
            if isinstance(obj,str):yield obj
            elif isinstance(obj,dict):
                for k,v in obj.items():yield k;yield from strings(v)
            elif isinstance(obj,list):
                for v in obj:yield from strings(v)
        for rel,p in files(repo).items():
            if p.suffix=='.json':values=strings(j(p))
            elif p.suffix=='.jsonl':values=(v for r in jl(p) for v in strings(r))
            elif p.suffix=='.md':values=[p.read_text(encoding='utf-8')]
            else:continue
            for value in values:
                # An official HTTP URL may legitimately contain /home/.
                without_urls=re.sub(r'https?://[^\s<>]+','',value)
                require(not pattern.search(without_urls),'Private path or credential pattern: '+rel)
        require(not any(p.suffix.lower() in {'.png','.jpg','.jpeg','.pdf','.zip'} for p in files(repo).values()),'Unexpected source image/archive')
    check('V20','No local absolute paths, credential patterns or image archives',check_privacy)
    def check_docs():
        for p in repo.rglob('*.md'):
            text=p.read_text(encoding='utf-8')
            for target in re.findall(r'(?<!!)\[[^\]]*\]\(([^)]+)\)',text):
                if re.match(r'https?://|mailto:|#',target):continue
                target=target.split('#',1)[0]
                require((p.parent/target).resolve().is_file(),'Broken link: '+p.name+' -> '+target)
        readme=(repo/'README.md').read_text(encoding='utf-8')
        require('3,000' in readme and '980' in readme and '人工确认 0' in readme and '图片生成 0' in readme)
        require('完整历史生成' in (repo/'docs/reproduction_scope.md').read_text(encoding='utf-8'))
    check('V21','Local documentation links and stated scope',check_docs)
    if check_hashes:check('V22','Exact release file set, sizes and hashes',lambda:require(verify(repo)['passed'],str(verify(repo)['problems'])))
    return {'passed':all(r['passed'] for r in results),'checks':len(results),'passed_checks':sum(r['passed'] for r in results),'failed_checks':sum(not r['passed'] for r in results),'results':results}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--repo',required=True);a=ap.parse_args()
    report=validate(a.repo);print(json.dumps(report,ensure_ascii=False,indent=2));return 0 if report['passed'] else 1

if __name__=='__main__':raise SystemExit(main())
